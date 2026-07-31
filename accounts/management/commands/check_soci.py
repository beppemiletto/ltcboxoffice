"""
Confronta il file anagrafica soci con gli utenti registrati sul sito.

Uso:
  python manage.py check_soci --file /path/to/soci.csv
  python manage.py check_soci --file /path/to/soci.xlsx
  python manage.py check_soci --file soci.csv --col-email Email --col-nome Nome --col-cognome Cognome

Colonne attese nel file (nomi default, modificabili con --col-*):
  Email, Nome, Cognome

Output:
  - Soci con account sul sito (pronti per is_socio=True)
  - Soci senza account (non ancora registrati)
  - Utenti sul sito non presenti in anagrafica soci
"""

import csv
import os
from django.core.management.base import BaseCommand, CommandError
from accounts.models import Account


def _load_file(path, col_email, col_nome, col_cognome):
    ext = os.path.splitext(path)[1].lower()
    rows = []

    if ext == '.csv':
        with open(path, encoding='utf-8-sig') as f:
            reader = csv.DictReader(f)
            for row in reader:
                rows.append(row)
    elif ext in ('.xlsx', '.xls', '.ods'):
        try:
            import openpyxl
            wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
            ws = wb.active
            headers = [str(c.value).strip() if c.value else '' for c in next(ws.iter_rows())]
            for ws_row in ws.iter_rows(min_row=2, values_only=True):
                rows.append({headers[i]: (str(v).strip() if v is not None else '') for i, v in enumerate(ws_row)})
        except ImportError:
            raise CommandError("openpyxl non installato. Installa con: pip install openpyxl")
    else:
        raise CommandError(f"Formato non supportato: {ext}. Usa .csv o .xlsx")

    soci = []
    for row in rows:
        email = row.get(col_email, '').strip().lower()
        nome = row.get(col_nome, '').strip()
        cognome = row.get(col_cognome, '').strip()
        if email:
            soci.append({'email': email, 'nome': nome, 'cognome': cognome, 'raw': row})
    return soci


class Command(BaseCommand):
    help = 'Confronta anagrafica soci con utenti del sito'

    def add_arguments(self, parser):
        parser.add_argument('--file', required=True, help='Percorso file CSV o Excel soci')
        parser.add_argument('--col-email', default='Email', help='Nome colonna email (default: Email)')
        parser.add_argument('--col-nome', default='Nome', help='Nome colonna nome (default: Nome)')
        parser.add_argument('--col-cognome', default='Cognome', help='Nome colonna cognome (default: Cognome)')
        parser.add_argument('--show-non-soci', action='store_true',
                            help='Mostra anche utenti del sito non presenti in anagrafica')

    def handle(self, *args, **options):
        path = options['file']
        if not os.path.exists(path):
            raise CommandError(f"File non trovato: {path}")

        soci = _load_file(path, options['col_email'], options['col_nome'], options['col_cognome'])
        self.stdout.write(f"Soci in anagrafica: {len(soci)}\n")

        all_accounts = {a.email.lower(): a for a in Account.objects.all()}
        soci_emails = {s['email'] for s in soci}

        # 1. Soci con account
        con_account = []
        senza_account = []
        for s in soci:
            acc = all_accounts.get(s['email'])
            if acc:
                con_account.append((s, acc))
            else:
                senza_account.append(s)

        # 2. Già soci sul sito
        gia_soci = [a for a in con_account if a[1].is_socio]
        da_attivare = [a for a in con_account if not a[1].is_socio]

        self.stdout.write('\n' + '='*60)
        self.stdout.write(self.style.SUCCESS(f'\n✓ SOCI CON ACCOUNT — già is_socio=True ({len(gia_soci)}):'))
        for s, acc in gia_soci:
            self.stdout.write(f'  {acc.email} — {acc.first_name} {acc.last_name}')

        self.stdout.write(self.style.WARNING(f'\n→ SOCI CON ACCOUNT — da attivare is_socio ({len(da_attivare)}):'))
        for s, acc in da_attivare:
            status = '(inattivo)' if not acc.is_active else ''
            self.stdout.write(f'  {acc.email} — {acc.first_name} {acc.last_name} {status}')

        self.stdout.write(self.style.ERROR(f'\n✗ SOCI SENZA ACCOUNT sul sito ({len(senza_account)}):'))
        for s in senza_account:
            self.stdout.write(f'  {s["email"]} — {s["nome"]} {s["cognome"]}')

        if options['show_non_soci']:
            non_soci = [a for e, a in all_accounts.items() if e not in soci_emails and not a.is_admin]
            self.stdout.write(f'\n  Utenti sul sito NON in anagrafica ({len(non_soci)}):')
            for a in sorted(non_soci, key=lambda x: x.last_name):
                self.stdout.write(f'  {a.email} — {a.first_name} {a.last_name}')

        self.stdout.write('\n' + '='*60)
        self.stdout.write(f'\nRiepilogo:')
        self.stdout.write(f'  Totale soci in anagrafica : {len(soci)}')
        self.stdout.write(f'  Con account sul sito      : {len(con_account)}')
        self.stdout.write(f'    già is_socio=True       : {len(gia_soci)}')
        self.stdout.write(f'    da attivare             : {len(da_attivare)}')
        self.stdout.write(f'  Senza account             : {len(senza_account)}')
        self.stdout.write('')
